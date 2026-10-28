# -*- coding: utf-8 -*-
"""Export an event record to Word (.docx) and Excel (.xlsx).

Usage:
    python -m event_build.export_office --event out/event.json --out exports/dir --title "My video"

Requires the optional ``office`` extra: ``pip install -e ".[office]"``.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

FONT = "Microsoft YaHei"


def _clock(seconds: float) -> str:
    total = max(0, int(round(seconds or 0)))
    return f"{total // 60:02d}:{total % 60:02d}"


def _load(event_path: Path, title: str | None) -> dict:
    event = json.loads(event_path.read_text(encoding="utf-8"))
    if title is None:
        review_path = event_path.parent / "review.json"
        if review_path.exists():
            try:
                title = json.loads(review_path.read_text(encoding="utf-8")).get("title")
            except Exception:
                title = None
    event["_title"] = title or event_path.parent.name
    return event


def _review(event: dict) -> list[dict]:
    review = event.get("review")
    if isinstance(review, dict) and isinstance(review.get("flags"), list):
        return review["flags"]
    return []


def _top_states(event: dict, limit: int = 8) -> list[dict]:
    states = event.get("states") or []
    return sorted(states, key=lambda s: -(s.get("attention", {}).get("emphasis", 0)))[:limit]


def _url_counts(event: dict) -> list[tuple[str, int]]:
    counts: dict[str, int] = {}
    for state in event.get("states") or []:
        url = state.get("url_hint") or "无 URL"
        counts[url] = counts.get(url, 0) + 1
    return sorted(counts.items(), key=lambda item: -item[1])


def _docx(event: dict, out: Path) -> Path:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    run = p.add_run(event["_title"])
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    meta = doc.add_paragraph(
        f"模式 {event.get('mode')} · 时长 {_clock(event.get('span_s', 0))} · "
        f"状态 {len(event.get('states') or [])} · 章节 {len(event.get('chapters') or [])} · "
        f"语元 {len(event.get('units') or [])} · 校验 "
        f"{'PASS' if (event.get('gates') or {}).get('all_ok') else 'FAIL'}"
    )
    meta.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    flags = _review(event)
    doc.add_heading("审阅标记", level=1)
    if not flags:
        doc.add_paragraph("无。")
    else:
        for flag in flags:
            doc.add_paragraph(f"[{flag.get('level')}] {flag.get('code')}: {flag.get('message')}", style="List Bullet")

    doc.add_heading("章节", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for index, header in enumerate(["时间", "依据", "URL", "停留", "注意力"]):
        table.rows[0].cells[index].text = header
    for chapter in event.get("chapters") or []:
        cells = table.add_row().cells
        cells[0].text = f"{_clock(chapter['t0'])}–{_clock(chapter['t1'])}"
        cells[1].text = "页面切换" if chapter.get("reason") == "url" else "话题"
        cells[2].text = chapter.get("url") or "—"
        cells[3].text = f"{round(chapter.get('dwell', 0))}s"
        cells[4].text = str(round(chapter.get("attention", 0)))

    doc.add_heading("注意力最高的状态", level=1)
    for state in _top_states(event):
        attention = state.get("attention") or {}
        label = state.get("label_hi") or state.get("label") or state.get("label_speech") or ""
        doc.add_paragraph(
            f"{_clock(state['t0'])}–{_clock(state['t1'])} · 注意力 {round(attention.get('emphasis', 0))} · "
            f"{state.get('url_hint') or '无 URL'} · {label[:80]}"
        )

    doc.add_heading("URL 分布", level=1)
    for url, count in _url_counts(event):
        doc.add_paragraph(f"{url} — {count} 个状态")

    path = out / "analysis.docx"
    doc.save(path)
    return path


def _xlsx(event: dict, out: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name=FONT, bold=True, color="FFFFFF")

    def header(ws, columns):
        ws.append(columns)
        for cell in ws[1]:
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "概览"
    header(ws, ["指标", "值"])
    gates = event.get("gates") or {}
    for key, value in [
        ("标题", event["_title"]),
        ("模式", event.get("mode")),
        ("时长", _clock(event.get("span_s", 0))),
        ("状态数", len(event.get("states") or [])),
        ("章节数", len(event.get("chapters") or [])),
        ("语元数", len(event.get("units") or [])),
        ("语音缺口", len(event.get("gaps") or [])),
        ("校验", "PASS" if gates.get("all_ok") else "FAIL"),
        ("审阅标记", len(_review(event))),
    ]:
        ws.append([key, value])
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 80

    chapters = wb.create_sheet("章节")
    header(chapters, ["开始", "结束", "开始(s)", "结束(s)", "依据", "URL", "停留(s)", "注意力"])
    for chapter in event.get("chapters") or []:
        chapters.append([
            _clock(chapter["t0"]), _clock(chapter["t1"]), round(chapter["t0"]), round(chapter["t1"]),
            "页面切换" if chapter.get("reason") == "url" else "话题",
            chapter.get("url") or "", round(chapter.get("dwell", 0)), round(chapter.get("attention", 0)),
        ])
    for column, width in {"A": 9, "B": 9, "C": 9, "D": 9, "E": 11, "F": 60, "G": 9, "H": 9}.items():
        chapters.column_dimensions[column].width = width

    review = wb.create_sheet("审阅标记")
    header(review, ["级别", "代码", "说明"])
    for flag in _review(event):
        review.append([flag.get("level"), flag.get("code"), flag.get("message")])
    review.column_dimensions["A"].width = 8
    review.column_dimensions["B"].width = 28
    review.column_dimensions["C"].width = 90

    attention = wb.create_sheet("注意力Top")
    header(attention, ["开始", "结束", "注意力", "URL", "标签"])
    for state in _top_states(event):
        att = state.get("attention") or {}
        attention.append([
            _clock(state["t0"]), _clock(state["t1"]), round(att.get("emphasis", 0)),
            state.get("url_hint") or "",
            (state.get("label_hi") or state.get("label") or state.get("label_speech") or "")[:120],
        ])
    attention.column_dimensions["A"].width = 9
    attention.column_dimensions["B"].width = 9
    attention.column_dimensions["C"].width = 10
    attention.column_dimensions["D"].width = 60
    attention.column_dimensions["E"].width = 70

    path = out / "analysis.xlsx"
    wb.save(path)
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description="Export an event record to Word and Excel")
    parser.add_argument("--event", required=True, help="path to event.json")
    parser.add_argument("--out", required=True, help="output directory")
    parser.add_argument("--title", default=None)
    parser.add_argument("--formats", default="docx,xlsx")
    args = parser.parse_args()

    event_path = Path(args.event).expanduser().resolve()
    out = Path(args.out).expanduser().resolve()
    out.mkdir(parents=True, exist_ok=True)
    event = _load(event_path, args.title)

    files: list[str] = []
    wanted = {item.strip() for item in args.formats.split(",") if item.strip()}
    if "docx" in wanted:
        files.append(str(_docx(event, out)))
    if "xlsx" in wanted:
        files.append(str(_xlsx(event, out)))
    print(json.dumps({"files": files}, ensure_ascii=False))


if __name__ == "__main__":
    main()
